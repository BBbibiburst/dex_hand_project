"""Update the supplied 2026-09-06 deck without changing its layout."""

from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

import cv2
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "张明远-2026.09.06.pptx"
OUTPUT = ROOT / "张明远-2026.09.19.pptx"
ASSETS = ROOT / "outputs" / "ppt_20260919_assets"
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"


def replace(shape, values: tuple[str, ...]) -> None:
    """Replace paragraph contents while retaining the template formatting."""
    paragraphs = shape.text_frame.paragraphs
    if len(values) > len(paragraphs):
        raise ValueError(f"{shape.name}: {len(values)} values for {len(paragraphs)} paragraphs")
    for index, paragraph in enumerate(paragraphs):
        value = values[index] if index < len(values) else ""
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        elif value:
            paragraph.add_run().text = value


def frame(video: Path, output: Path, ratio: float = 0.60) -> Image.Image:
    capture = cv2.VideoCapture(str(video))
    count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
    capture.set(cv2.CAP_PROP_POS_FRAMES, max(0, int(count * ratio)))
    ok, image = capture.read()
    capture.release()
    if not ok:
        raise RuntimeError(f"Could not read {video}")
    image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
    result = Image.fromarray(image)
    result.save(output)
    return result


def cover(im: Image.Image, size: tuple[int, int]) -> Image.Image:
    scale = max(size[0] / im.width, size[1] / im.height)
    resized = im.resize((int(im.width * scale), int(im.height * scale)), Image.Resampling.LANCZOS)
    left = (resized.width - size[0]) // 2
    top = (resized.height - size[1]) // 2
    return resized.crop((left, top, left + size[0], top + size[1]))


def label_panel(im: Image.Image, title: str, metric: str, color: str) -> Image.Image:
    out = cover(im, (500, 780))
    overlay = Image.new("RGBA", out.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    draw.rectangle((0, 0, 500, 122), fill=(7, 34, 78, 218))
    draw.rectangle((0, 715, 500, 780), fill=color)
    draw.text((28, 26), title, font=ImageFont.truetype(FONT_BOLD, 35), fill="white")
    draw.text((28, 75), metric, font=ImageFont.truetype(FONT, 24), fill=(211, 228, 248))
    return Image.alpha_composite(out.convert("RGBA"), overlay).convert("RGB")


def make_visuals() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    videos = {
        "stack": ROOT / "outputs/demo_videos/stack_egad_E0_v18_diagnostic.mp4",
        # This is the verified four-finger, palm-side tactile-contact demo.
        # Do not use top100_push.mp4 here: that recording shows the obsolete
        # dorsal/back-of-hand contact orientation.
        "push": ROOT / "outputs/demo_videos/push_four_finger_v32.mp4",
        "nut": ROOT / "outputs/demo_videos/nut_assembly_v41.mp4",
    }
    fallback = {
        "stack": ROOT / "outputs/demo_videos/top100_stack.mp4",
        "push": ROOT / "outputs/archive/2026-09-07-experiments/push_debug_palm_face_v9/c01.mp4",
        "nut": ROOT / "outputs/demo_videos/top100_nut_assembly.mp4",
    }
    stills: dict[str, Path] = {}
    for name, video in videos.items():
        source = video if video.is_file() else fallback[name]
        output = ASSETS / f"{name}_latest.png"
        frame(source, output)
        stills[name] = output

    panels = [
        label_panel(Image.open(stills["stack"]), "Stack", "Task-valid 44 / 48", "#e39a2c"),
        label_panel(Image.open(stills["push"]), "Push", "Task-valid 71 / 76", "#1674c7"),
        label_panel(Image.open(stills["nut"]), "NutAssembly", "100 / 100 + stability", "#20985f"),
    ]
    collage = Image.new("RGB", (1500, 920), "white")
    for index, panel in enumerate(panels):
        collage.paste(panel, (500 * index, 0))
    draw = ImageDraw.Draw(collage)
    draw.rectangle((0, 780, 1500, 920), fill="#092f70")
    draw.text(
        (750, 833),
        "TASK-VALID BENCHMARK · C MUJOCO VERIFICATION",
        font=ImageFont.truetype(FONT_BOLD, 35),
        fill="white",
        anchor="mm",
    )
    collage_path = ASSETS / "three_tasks_latest.png"
    collage.save(collage_path)

    overview = ROOT / "outputs/nut_assembly_top100_final/nut_variant_overview.png"
    return {
        "image4.png": collage_path,
        "image5.png": stills["stack"],
        "image6.png": stills["push"],
        "image7.png": stills["push"],
        "image8.png": overview,
        "image9.png": stills["nut"],
        "media1.mp4": videos["stack"] if videos["stack"].is_file() else fallback["stack"],
        "media2.mp4": videos["push"] if videos["push"].is_file() else fallback["push"],
        "media3.mp4": videos["nut"] if videos["nut"].is_file() else fallback["nut"],
    }


def update_text() -> None:
    prs = Presentation(TEMPLATE)
    # Map: slide number -> shape number -> replacement paragraphs.
    content: dict[int, dict[int, tuple[str, ...]]] = {
        1: {2: ("2026年09月19日 | 组会进展汇报",)},
        2: {
            2: ("进展一：五任务评测域重构与最新结果",),
            3: ("：本周总体进展",),
            4: ("抬升 Lift", "100 / 100"),
            5: ("搬运放置", "PickPlace 100 / 100"),
            6: ("堆叠 Stack", "44 / 48"),
            7: ("推物 Push", "71 / 76"),
            8: ("螺母装配", "100 / 100"),
            9: (
                "面向 RM75B + 6 驱动 Dex Hand，完成五任务评测口径重构。",
                "满分任务：",
                "Lift、PickPlace、NutAssembly 均达到 100 / 100。",
                "",
                "Task-valid 主指标：",
                "Stack 44 / 48（91.7%）；Push 71 / 76（93.4%）。",
                "完整 Top100 历史覆盖分别为 95 / 100、91 / 100。",
                "当前重点：",
                "解决 Stack 4 个与 Push 5 个真实算法难例。",
            ),
        },
        3: {
            0: (
                "上阶段问题：",
                "五个任务强行共用 Top100，混淆任务不适配与算法失败。",
                "本周完成：",
                "建立 Stack 48 与 Push 76 个 Task-valid 对象集，并统一五任务报告。",
                "原则：",
                "筛选只依赖几何与独立静置结果，不读取策略执行成败。",
                "当前瓶颈已收敛为 Stack 4 个、Push 5 个难例。",
            ),
            3: ("进展一：五任务评测域重构与最新结果",),
            4: ("：上周问题与本周结果",),
            5: ("同一 Top100", "口径不合理"),
            6: ("2 个任务域", "本周重构"),
            7: ("5 任务", "统一报告"),
            8: ("物理定义", "→", "事前筛选"),
            9: ("任务执行", "→", "稳定复验"),
            10: ("问题收敛：从评测口径修正到 9 个真实算法难例",),
        },
        4: {
            2: ("进展一：五任务评测域重构与最新结果",),
            3: ("：Task-valid 生成与统一报告流程",),
            4: ("Top100", "输入"),
            5: ("几何 + 静置", "预检"),
            6: ("Task-valid",),
            8: ("C MuJoCo", "执行"),
            9: ("稳定性", "复验"),
            10: ("五任务", "报告"),
            11: ("版本化清单 + 排除原因 + 双层成功率",),
            12: (
                "统一输入：",
                "Top100 完整保留，",
                "Stack / Push 通过事前物理规则生成任务子集。",
                "统一产物：",
                "版本化 domain JSON、逐项测量值、排除原因和 summary.json。",
                "统一指标：",
                "主指标为 Task-valid success，辅指标为 Full-catalogue coverage。",
                "工程能力：",
                "清单可重建、批次可续跑、结果可审计。",
            ),
        },
        5: {2: ("进展一：五任务评测域重构与最新结果",), 3: ("：最新代表性执行画面",)},
        6: {
            0: (
                "评测域：",
                "Top100 中筛得 48 个具有稳定支撑几何的对象。",
                "运输约束：",
                "真实悬空；transport 阶段逐帧审计平台接触。",
                "放置约束：",
                "进入落点后冻结 XY，仅允许垂直下降。",
                "稳定验证：",
                "承重释放、机械手退让，再连续复验 150 帧。",
                "",
            ),
            3: ("进展二：Stack Task-valid 评测",),
            4: ("：物理有效域与执行约束",),
            5: ("Lift", "抓持前缀"),
            7: ("悬空", "横向运输"),
            9: ("冻结 XY", "垂直下降"),
            11: ("承重", "释放复验"),
            12: ("硬约束", "运输零平台接触 · 退让后稳定"),
        },
        7: {
            2: ("进展二：Stack Task-valid 评测",),
            3: ("：最新结果与严格成功判据",),
            4: ("Task-valid", "48"),
            5: ("成功", "44"),
            6: ("剩余", "4"),
            11: ("Task-valid 91.7% · Top100 coverage 95%",),
            12: (
                "Task-valid 域：",
                "48 个对象通过几何与独立静置预检，52 个按规则排除。",
                "当前成功率：",
                "历史物理执行与新评测域交集为 44 / 48（91.7%）。",
                "完整目录覆盖：",
                "原 Top100 历史成功仍保留为 95 / 100 辅助指标。",
                "结果意义：",
                "剩余 4 个均为有效对象，必须继续解决而不能删除。",
                "",
            ),
        },
        8: {
            2: ("进展二：Stack Task-valid 评测",),
            3: ("：剩余 4 个难例诊断",),
            4: ("EGAD E0",),
            6: ("玩具冰箱",),
            8: ("尿布桶",),
            10: ("M40",),
            12: ("平台接触审计",),
            14: ("抓持丢失",),
            15: ("主线变化：前三项已通过稳定复验，失败集中于运输审计；M40 集中于抓持保持",),
            16: (
                "当前结果：",
                "Stack Task-valid = 44 / 48，仍有 4 个有效难例。",
                "前三个对象：",
                "EGAD E0、玩具冰箱、尿布桶可稳定放置，但早期接触平台。",
                "第四个对象：",
                "染发凝胶 M40 在运输阶段丢失，属于抓持稳定域不足。",
                "修复方向：",
                "v24 调整跨平台路径；M40 定向搜索抓持姿态。",
            ),
        },
        9: {1: ("进展二：Stack Task-valid 评测",)},
        10: {
            2: ("进展三：Push 四指掌侧触觉面闭环",),
            3: ("：Task-valid 结果与剩余难例",),
            4: ("Push Task-valid", "71 / 76"),
            5: ("任务有效域", "76"),
            6: ("当前成功", "71"),
            7: ("距离类", "2"),
            8: ("接触类", "2"),
            9: ("制动类", "1"),
            10: (
                "任务定义：",
                "四指掌侧触觉面准静态滑动；排除球形滚动和接触跨度不足对象。",
                "最新结果：",
                "Task-valid 成功 71 / 76（93.4%）；Top100 历史覆盖 91 / 100。",
                "剩余对象：",
                "EGAD G0、布丁盒、Reebok 鞋、手电筒和 Jerky Tonic。",
                "失败类别：",
                "目标距离不足、手—桌接触、接触丢失与近目标制动不足。",
                "当前修复：",
                "v31 已加入 near-target deceleration，待完整物理回归。",
            ),
        },
        11: {2: ("进展三：Push 四指掌侧触觉面闭环",), 3: ("：最新代表性闭环推动执行",)},
        12: {
            2: ("进展四：NutAssembly 程序化装配",),
            3: ("：100 个变体完整保留",),
            0: (
                "任务设定：",
                "固定螺母 + 自由螺栓，欠驱动手完成抓持、对孔与插入。",
                "评测域：",
                "100 项是程序化插销—插孔组合，不是日常物体 Top100。",
                "变体空间：",
                "10 种截面 × 10 种尺寸族，覆盖圆形及 3–16 边形。",
                "结论：",
                "执行 100 / 100，独立稳定性回放 100 / 100，不删变体。",
            ),
        },
        13: {1: ("进展四：NutAssembly 程序化装配",)},
        14: {2: ("进展四：NutAssembly 程序化装配",)},
        15: {
            3: ("进展四：NutAssembly 程序化装配",),
            4: ("：执行与独立稳定性复验结果",),
            5: ("NutAssembly 变体", "100 / 100"),
            6: ("100", "执行成功"),
            7: ("100", "稳定回放"),
            8: ("全部变体通过同一物理成功判据与独立终态复验",),
            1: (
                "变体组成：",
                "10 种截面族 × 10 种尺寸族，覆盖圆、方及规则多边形。",
                "执行结果：",
                "v41 在 C MuJoCo 中完成全部 100 条完整插入轨迹。",
                "独立复验：",
                "100 个终态重新加载后全部通过持续稳定性回放。",
                "结论：",
                "该程序化域完整保留；可另设 Core 子集用于快速回归。",
            ),
        },
        16: {
            0: (
                "本周已完成：",
                "建立 Stack / Push Task-valid 清单与五任务统一报告口径。",
                "Stack 当前状态：",
                "44 / 48；前三项为平台接触审计，M40 为运输中丢失。",
                "Push 当前状态：",
                "71 / 76；分为距离、桌面接触、接触保持与制动问题。",
                "工程保障：",
                "公共 batch workflow 已抽取，155 项测试与 Ruff 检查通过。",
                "下一目标：",
                "完成 9 个难例单对象回归，再运行正式 Task-valid 全量批次。",
                "保持物理成功与稳定复验标准不降低。",
            ),
            2: ("进展五：当前状态与后续推进",),
            3: ("：五任务最新收敛情况",),
            4: ("Lift", "100%"),
            6: ("PickPlace", "100%"),
            8: ("Stack", "91.7%"),
            9: ("Push", "93.4%"),
            10: ("NutAssembly", "100%"),
            11: ("当前重点：Stack 4 个 + Push 5 个 Task-valid 难例定向修复",),
        },
        17: {1: ("进展四：NutAssembly 程序化装配",), 2: ("：最新代表性插入与稳定驻留",)},
        18: {3: ("哈尔滨工业大学", "2026年09月19日")},
    }
    for slide_number, shapes in content.items():
        for shape_number, values in shapes.items():
            replace(prs.slides[slide_number - 1].shapes[shape_number], values)
    prs.core_properties.title = "2026年09月19日组会进展汇报"
    prs.save(OUTPUT)


def replace_package_media(media: dict[str, Path]) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_root = Path(tmp)
        expanded = tmp_root / "expanded"
        with zipfile.ZipFile(OUTPUT) as archive:
            archive.extractall(expanded)
        for name, source in media.items():
            target = expanded / "ppt" / "media" / name
            if not source.is_file():
                raise FileNotFoundError(source)
            shutil.copyfile(source, target)
        rebuilt = tmp_root / "rebuilt.pptx"
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in expanded.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(expanded))
        shutil.copyfile(rebuilt, OUTPUT)


def main() -> None:
    media = make_visuals()
    update_text()
    replace_package_media(media)
    print(f"saved={OUTPUT} slides=18 updated_media={len(media)}")


if __name__ == "__main__":
    main()
