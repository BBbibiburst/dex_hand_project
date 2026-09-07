"""Build the 2026-09-06 group-meeting deck from current formal results."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "张明远-2026.09.06.pptx"
ASSETS = ROOT / "outputs" / "ppt_20260906_assets"

NAVY = RGBColor(23, 58, 138)
BLUE = RGBColor(38, 103, 190)
CYAN = RGBColor(29, 161, 181)
GREEN = RGBColor(27, 111, 42)
ORANGE = RGBColor(181, 104, 21)
RED = RGBColor(169, 62, 62)
DARK = RGBColor(25, 27, 31)
GRAY = RGBColor(92, 99, 112)
LIGHT = RGBColor(244, 247, 251)
PALE_BLUE = RGBColor(234, 241, 255)
PALE_GREEN = RGBColor(234, 247, 238)
PALE_ORANGE = RGBColor(255, 244, 228)
WHITE = RGBColor(255, 255, 255)
FONT = "微软雅黑"


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=18,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, *, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = fill if line is None else line
    return sh


def header(slide, section, title, subtitle=None):
    add_text(slide, "研究进展", 0.55, 0.24, 1.3, 0.3, size=11, color=GRAY, bold=True)
    add_text(slide, section, 1.55, 0.24, 5.2, 0.3, size=11, color=NAVY, bold=True)
    add_text(slide, title, 0.55, 0.68, 12.15, 0.52, size=26, color=DARK, bold=True)
    rect(slide, 0.55, 1.25, 0.75, 0.055, NAVY, radius=False)
    if subtitle:
        add_text(slide, subtitle, 1.45, 1.16, 11.2, 0.3, size=11, color=GRAY)


def footer(slide, number):
    add_text(slide, "张明远 · 2026.09.06", 0.55, 7.15, 2.8, 0.18, size=8, color=GRAY)
    add_text(slide, str(number), 12.25, 7.12, 0.45, 0.2, size=8, color=GRAY, align=PP_ALIGN.RIGHT)


def bullet_box(slide, items, x, y, w, h, *, size=16, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.text = "• " + p.text
    return box


def metric(slide, x, y, w, value, label, fill, color):
    rect(slide, x, y, w, 1.12, fill)
    add_text(
        slide,
        value,
        x + 0.12,
        y + 0.13,
        w - 0.24,
        0.46,
        size=27,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide, label, x + 0.12, y + 0.66, w - 0.24, 0.25, size=11, color=DARK, align=PP_ALIGN.CENTER
    )


def add_image(slide, path, x, y, w, h):
    rect(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, RGBColor(225, 230, 238), radius=False)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def flow(slide, labels, y, colors=None):
    colors = colors or [PALE_BLUE] * len(labels)
    gap, left, total = 0.18, 0.65, 12.0
    width = (total - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        x = left + i * (width + gap)
        rect(slide, x, y, width, 0.74, colors[i])
        add_text(
            slide,
            label,
            x + 0.08,
            y + 0.12,
            width - 0.16,
            0.42,
            size=13,
            color=DARK,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if i < len(labels) - 1:
            add_text(
                slide,
                "→",
                x + width,
                y + 0.19,
                gap,
                0.3,
                size=16,
                color=NAVY,
                bold=True,
                align=PP_ALIGN.CENTER,
            )


def load_summary(name):
    return json.loads((ROOT / name).read_text(encoding="utf-8"))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    # 1 cover
    s = prs.slides.add_slide(blank)
    slides.append(s)
    rect(s, 0, 0, 13.333, 7.5, WHITE, radius=False)
    rect(s, 0, 0, 0.22, 7.5, NAVY, radius=False)
    rect(s, 0.7, 1.25, 1.05, 0.09, CYAN, radius=False)
    add_text(
        s, "2026年09月06日  组会进展汇报", 0.7, 0.62, 5.2, 0.45, size=17, color=NAVY, bold=True
    )
    add_text(
        s,
        "多任务仿真示教迁移\nTop100 全量验证进展",
        0.7,
        1.62,
        9.2,
        1.45,
        size=34,
        color=DARK,
        bold=True,
    )
    add_text(s, "Stack · Push · NutAssembly", 0.72, 3.35, 6.5, 0.45, size=20, color=CYAN, bold=True)
    add_text(s, "张明远", 0.72, 5.65, 2.5, 0.45, size=21, color=DARK, bold=True)
    add_text(s, "哈尔滨工业大学", 0.72, 6.15, 3.0, 0.3, size=13, color=GRAY)
    for i, (v, c) in enumerate((("100", GREEN), ("100", GREEN), ("95", ORANGE))):
        rect(s, 9.35 + i * 1.15, 4.55 - i * 0.38, 0.9, 1.25 + i * 0.38, c, radius=False)
        add_text(
            s,
            v,
            9.35 + i * 1.15,
            4.83 - i * 0.38,
            0.9,
            0.4,
            size=20,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    # 2 overview
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展一：全量结果", "本周结论")
    metric(s, 0.65, 1.65, 2.25, "100 / 100", "Lift（既有基线）", PALE_GREEN, GREEN)
    metric(s, 3.05, 1.65, 2.25, "100 / 100", "PickPlace（已验证）", PALE_GREEN, GREEN)
    metric(s, 5.45, 1.65, 2.25, "95 / 100", "Stack v18", PALE_ORANGE, ORANGE)
    metric(s, 7.85, 1.65, 2.25, "100 / 100", "Push v8", PALE_GREEN, GREEN)
    metric(s, 10.25, 1.65, 2.25, "100 / 100", "NutAssembly v22", PALE_GREEN, GREEN)
    add_text(s, "本周新增完成", 0.65, 3.18, 2.1, 0.35, size=16, color=NAVY, bold=True)
    bullet_box(
        s,
        [
            "Push：掌心皮肤触觉闭环，Top100 全量 100%",
            "NutAssembly：固定螺母 + 欠驱动抓持插入，100 个程序化变体全通过",
            "Stack：严格悬空搬运、垂直放置与驻留复验，提升至 95%",
        ],
        0.8,
        3.62,
        11.8,
        2.0,
        size=17,
    )
    rect(s, 0.65, 6.05, 12.0, 0.55, PALE_BLUE)
    add_text(
        s,
        "当前瓶颈已收敛到 Stack 的 5 个极端几何对象，其余两项已闭环。",
        0.85,
        6.17,
        11.6,
        0.25,
        size=15,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 3 unified pipeline
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展一：全量结果", "统一任务生成与验证流程")
    flow(
        s,
        ["Top100 物体 / 变体", "任务候选搜索", "C MuJoCo 执行", "稳定性驻留复验", "轨迹与报告落盘"],
        1.75,
        [PALE_BLUE, PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_GREEN],
    )
    add_text(s, "三项任务共享的工程约束", 0.68, 3.0, 3.5, 0.4, size=18, color=NAVY, bold=True)
    bullet_box(
        s,
        [
            "逐对象日志、断点续跑、并行批处理",
            "轨迹记录 qpos / qvel / ctrl / stage / contact",
            "最终判定由 C MuJoCo 完整物理回放给出",
            "成功轨迹必须通过任务特定的持续稳定性条件",
        ],
        0.85,
        3.48,
        5.6,
        2.5,
        size=16,
    )
    rect(s, 7.0, 3.15, 5.45, 2.7, LIGHT)
    add_text(
        s,
        "统一的成功定义",
        7.3,
        3.45,
        4.85,
        0.35,
        size=19,
        color=DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "可执行  +  物理稳定  +  可复验",
        7.3,
        4.15,
        4.85,
        0.5,
        size=23,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "避免仅凭末帧位置或运动学条件判成功",
        7.3,
        5.05,
        4.85,
        0.3,
        size=13,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )

    # 4 stack method
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展二：Stack", "无斜坡悬空搬运与垂直放置")
    add_image(s, ASSETS / "stack_mid.png", 0.65, 1.55, 5.6, 3.15)
    flow(
        s,
        ["复用 Lift 抓持", "悬空横移", "冻结 XY", "垂直下降", "承重释放"],
        5.05,
        [PALE_BLUE, PALE_ORANGE, PALE_BLUE, PALE_BLUE, PALE_GREEN],
    )
    bullet_box(
        s,
        [
            "固定、焊接的实体平台；不使用斜坡或隐式支撑",
            "平台沿抓持/搬运方向搜索，降低腕部不可达与侧碰",
            "进入落点阈值后冻结水平坐标，保证纯垂直下降",
            "运输阶段逐帧审计平台接触：必须全程 suspended_transport=True",
        ],
        6.65,
        1.7,
        5.95,
        3.05,
        size=16,
    )
    rect(s, 6.72, 5.35, 5.65, 0.72, PALE_GREEN)
    add_text(
        s,
        "释放后退 100 帧 + 稳定复验 150 帧",
        6.9,
        5.55,
        5.3,
        0.25,
        size=16,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 5 stack result
    stack = load_summary("outputs/stack_top100_v18/summary.json")
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展二：Stack", "Top100 正式批次：95 / 100")
    metric(s, 0.75, 1.55, 3.0, "95%", "C MuJoCo + 驻留复验", PALE_ORANGE, ORANGE)
    metric(s, 4.0, 1.55, 3.0, "95", "成功对象", PALE_GREEN, GREEN)
    metric(s, 7.25, 1.55, 3.0, "5", "剩余失败", RGBColor(250, 232, 232), RED)
    metric(s, 10.5, 1.55, 2.0, "v18", "流程版本", PALE_BLUE, NAVY)
    add_text(s, "5 个剩余对象", 0.75, 3.1, 3.0, 0.35, size=18, color=RED, bold=True)
    failures = [
        r["object_id"].split(":", 1)[-1] for r in stack["results"] if r["status"] != "SUCCESS"
    ]
    bullet_box(s, failures, 0.9, 3.52, 6.2, 2.75, size=14)
    rect(s, 7.45, 3.25, 5.0, 2.55, LIGHT)
    add_text(
        s,
        "共同失败原因",
        7.75,
        3.55,
        4.4,
        0.35,
        size=18,
        color=DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "object was not stacked\nstably on the support",
        7.75,
        4.22,
        4.4,
        0.85,
        size=21,
        color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        "说明抓取与运输已不再是主导瓶颈，问题集中于落点/释放后的姿态稳定。",
        7.75,
        5.28,
        4.4,
        0.45,
        size=13,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )

    # 6 stack diagnosis
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展二：Stack", "最后 5 个对象：问题定位与修正方向")
    cols = [
        ("几何特征", ["薄片 / 长条", "非对称质量分布", "小支撑多边形"], PALE_BLUE, NAVY),
        ("失败机制", ["落地瞬间翻滚", "释放弹性回弹", "稳定域过窄"], RGBColor(250, 232, 232), RED),
        (
            "下一步",
            ["按几何估计稳定落姿", "平台接触反馈微调", "释放速度与时序搜索"],
            PALE_GREEN,
            GREEN,
        ),
    ]
    for i, (title, items, fill, c) in enumerate(cols):
        x = 0.65 + i * 4.15
        rect(s, x, 1.62, 3.8, 3.9, fill)
        add_text(
            s, title, x + 0.2, 1.92, 3.4, 0.4, size=20, color=c, bold=True, align=PP_ALIGN.CENTER
        )
        bullet_box(s, items, x + 0.35, 2.65, 3.1, 2.05, size=16)
    rect(s, 0.8, 5.92, 11.7, 0.62, PALE_ORANGE)
    add_text(
        s,
        "目标：保持“无斜坡 + 真悬空 + 纯垂直放置”约束不变，只优化落姿与释放。",
        1.0,
        6.08,
        11.3,
        0.28,
        size=16,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 7 push method
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展三：Push", "掌心皮肤触觉闭环推动")
    add_image(s, ASSETS / "push_mid.png", 0.65, 1.55, 5.6, 3.15)
    bullet_box(
        s,
        [
            "以掌心皮肤 taxel 为主接触面，降低指尖几何差异影响",
            "接触丢失时根据触觉、接触距离与目标距离趋势联合重捕获",
            "推进中自适应修正航向，避免圆柱/球体滚偏",
            "终点使用反馈式微推，并连续验证目标区域稳定性",
        ],
        6.65,
        1.7,
        5.9,
        3.0,
        size=16,
    )
    flow(
        s,
        ["接近", "触觉建联", "闭环推进", "丢失重捕获", "目标复验"],
        5.22,
        [PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_ORANGE, PALE_GREEN],
    )

    # 8 push result
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展三：Push", "Top100 全量结果：100 / 100")
    metric(s, 0.75, 1.55, 3.15, "100%", "正式成功率", PALE_GREEN, GREEN)
    metric(s, 4.15, 1.55, 3.15, "100", "成功对象", PALE_GREEN, GREEN)
    metric(s, 7.55, 1.55, 2.15, "0", "失败对象", PALE_BLUE, NAVY)
    metric(s, 9.95, 1.55, 2.55, "v8", "最终流程版本", PALE_BLUE, NAVY)
    add_text(s, "关键收敛点", 0.8, 3.15, 2.5, 0.35, size=18, color=NAVY, bold=True)
    bullet_box(
        s,
        [
            "掌心有效触觉覆盖扩展至边缘，消除接触盲区",
            "推动高度降低，减少物体倾覆力矩",
            "将“是否接触”升级为距离趋势 + 触觉 + 任务误差的联合反馈",
            "对难例采用自适应转向与两阶段末端修正",
        ],
        0.95,
        3.62,
        11.4,
        2.35,
        size=17,
    )
    rect(s, 0.8, 6.15, 11.7, 0.48, PALE_GREEN)
    add_text(
        s,
        "Push 已达到任务目标，可直接进入数据导出与演示整理。",
        1.0,
        6.25,
        11.3,
        0.25,
        size=15,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 9 nut method
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展四：NutAssembly", "固定螺母 + 欠驱动手抓持螺栓插入")
    add_image(s, ASSETS / "nut_mid.png", 0.65, 1.55, 5.6, 3.15)
    bullet_box(
        s,
        [
            "螺母固定在桌面，保留真实壁面碰撞与形状配合",
            "欠驱动六执行器手抓持带把手螺栓，完成搬运与插入",
            "100 个可复现程序化变体：圆形、方形及多边形截面",
            "成功需满足位置、竖直度、插入深度、速度与持续稳定帧",
        ],
        6.65,
        1.7,
        5.9,
        3.0,
        size=16,
    )
    flow(
        s,
        ["专用抓持", "抬升", "对孔运输", "竖直插入", "长驻留"],
        5.22,
        [PALE_BLUE, PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_GREEN],
    )

    # 10 nut result
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展四：NutAssembly", "100 个程序化变体：100 / 100")
    metric(s, 0.75, 1.55, 3.15, "100%", "正式成功率", PALE_GREEN, GREEN)
    metric(s, 4.15, 1.55, 3.15, "100", "形状/尺寸变体", PALE_GREEN, GREEN)
    metric(s, 7.55, 1.55, 2.25, "20", "连续稳定帧", PALE_BLUE, NAVY)
    metric(s, 10.05, 1.55, 2.45, "v22", "最终流程版本", PALE_BLUE, NAVY)
    add_text(s, "变体空间", 0.8, 3.2, 2.2, 0.35, size=18, color=NAVY, bold=True)
    flow(
        s,
        ["圆形", "方形", "三角形", "5–16 边形", "10 个尺寸族"],
        3.72,
        [PALE_BLUE, PALE_BLUE, PALE_ORANGE, PALE_BLUE, PALE_GREEN],
    )
    bullet_box(
        s,
        [
            "三角棱柱对外接半径加上限，保持公共抓持技能可迁移",
            "对最慢收敛变体启用更长 settle，仍使用同一物理成功判据",
            "结果说明欠驱动手在较大被动角度误差下仍可依靠导向配合完成插入",
        ],
        0.95,
        5.05,
        11.4,
        1.25,
        size=15,
    )

    # 11 verification
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "进展五：验证与产物", "成功不是末帧截图：任务特定的严格复验")
    rows = [
        ("Stack", "运输零平台接触；释放后退；150 帧平台稳定", ORANGE),
        ("Push", "进入目标区域；闭环停止；连续目标稳定复验", CYAN),
        ("NutAssembly", "位置 + 竖直度 + 深度 + 低速；连续 20 帧", GREEN),
    ]
    for i, (name, rule, c) in enumerate(rows):
        y = 1.62 + i * 1.2
        rect(s, 0.75, y, 2.0, 0.8, c)
        add_text(
            s,
            name,
            0.85,
            y + 0.18,
            1.8,
            0.3,
            size=17,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        rect(s, 2.95, y, 9.55, 0.8, LIGHT)
        add_text(s, rule, 3.25, y + 0.17, 8.95, 0.32, size=16, color=DARK, bold=True)
    add_text(s, "本周新增/完善的可交付产物", 0.78, 5.35, 4.0, 0.35, size=18, color=NAVY, bold=True)
    bullet_box(
        s,
        [
            "三项任务独立生成器与 Top100 批处理入口",
            "逐对象 manifest / episode.npz / summary.json",
            "任务通用视频回放工具与代表性 MP4 演示",
            "失败对象保留完整阶段与接触诊断",
        ],
        0.95,
        5.78,
        11.3,
        1.0,
        size=14,
    )

    # 12 next
    s = prs.slides.add_slide(blank)
    slides.append(s)
    header(s, "下一步", "Stack 5 个难例收尾与统一发布")
    steps = [
        ("1", "落姿生成", "基于支撑多边形与质心估计稳定姿态"),
        ("2", "接触微调", "下降末端利用平台接触反馈修正 XY / yaw"),
        ("3", "释放搜索", "对释放速度、卸载时序与退让方向做小规模搜索"),
        ("4", "全量复验", "重新跑 Top100，目标 100/100 后统一导出"),
        ("5", "整理发布", "演示、测试、文档、提交、标签与推送"),
    ]
    for i, (n, t, d) in enumerate(steps):
        y = 1.48 + i * 1.02
        rect(s, 0.75, y, 0.62, 0.62, NAVY)
        add_text(
            s, n, 0.75, y + 0.12, 0.62, 0.25, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER
        )
        add_text(s, t, 1.65, y + 0.05, 2.0, 0.3, size=17, color=DARK, bold=True)
        add_text(s, d, 3.55, y + 0.06, 8.6, 0.3, size=15, color=GRAY)
    rect(s, 0.75, 6.65, 11.75, 0.42, PALE_ORANGE)
    add_text(
        s,
        "下周硬目标：Stack 100/100；三项任务统一 C MuJoCo 稳定性复验通过。",
        0.95,
        6.72,
        11.35,
        0.25,
        size=15,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 13 closing
    s = prs.slides.add_slide(blank)
    slides.append(s)
    rect(s, 0, 0, 13.333, 7.5, WHITE, radius=False)
    rect(s, 0, 0, 0.22, 7.5, NAVY, radius=False)
    add_text(s, "请批评指正！", 0.75, 2.2, 7.0, 0.8, size=38, color=DARK, bold=True)
    add_text(
        s,
        "Push 100/100 · NutAssembly 100/100 · Stack 95/100",
        0.78,
        3.35,
        9.5,
        0.45,
        size=20,
        color=NAVY,
        bold=True,
    )
    add_text(
        s, "张明远  |  哈尔滨工业大学  |  2026年09月06日", 0.78, 5.7, 7.8, 0.35, size=14, color=GRAY
    )
    rect(s, 9.55, 1.55, 2.7, 4.4, NAVY)
    add_text(
        s,
        "100\n100\n95",
        9.55,
        2.0,
        2.7,
        2.9,
        size=38,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    for i, slide in enumerate(slides[1:-1], 2):
        footer(slide, i)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
